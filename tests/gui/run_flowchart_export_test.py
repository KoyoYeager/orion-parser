"""Test flowchart export from GUI — every format."""

import os
import sys

os.environ["QT_QPA_PLATFORM"] = "offscreen"

from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

app = QApplication([])

from orionparser.gui.main_window import MainWindow
from orionparser.registry import get_pipeline
from orionparser.core.encoding import read_file

SS = Path(__file__).parent / "screenshots" / "flowchart_export"
SS.mkdir(parents=True, exist_ok=True)

issues: list[str] = []


def check(name, cond, msg=""):
    if not cond:
        issues.append(f"{name}: {msg}")
        print(f"  FAIL: {name} - {msg}")
    else:
        print(f"  OK: {name}")


# Setup
p = Path("C:/tmp/fc_test/processor.py")
result = get_pipeline(p).analyze_file(p)
source = read_file(p)

win = MainWindow()
win.show()
QTest.qWaitForWindowExposed(win)
win._on_file_loaded(result, source, str(p))
QTest.qWait(300)

# Switch to graph mode + flowchart
win._switch_mode(1)
QTest.qWait(200)

selector = win._graph_view._selector
for i in range(selector.count()):
    item = selector.item(i)
    if item and item.data(Qt.ItemDataRole.UserRole) == "flowchart":
        selector.setCurrentRow(i)
        QTest.qWait(300)
        break

# Select process_data function
combo = win._graph_view._focus_combo
for i in range(combo.count()):
    if combo.itemData(i) == "process_data":
        combo.setCurrentIndex(i)
        win._graph_view._on_focus_selected(i)
        QTest.qWait(300)
        break

canvas = win._graph_view.canvas
print(f"Canvas nodes: {len(canvas._node_items)}")
check("has_nodes", len(canvas._node_items) > 5, f"only {len(canvas._node_items)}")

# Test every export format
print("\n=== Export Tests ===")

_FORMATS = {
    "png": "PNG",
    "svg": "SVG",
    "pdf": "PDF",
    "jpg": "JPEG",
    "dot": "DOT",
    "xlsx": "Excel",
    "html": "HTML",
    "json": "JSON",
    "mermaid": "Mermaid",
    "plantuml": "PlantUML",
    "drawio": "Draw.io",
    "xml": "XML",
    "yaml": "YAML",
}

_EXT_MAP = {"mermaid": "md", "plantuml": "puml", "drawio": "drawio"}

for fmt, name in _FORMATS.items():
    ext = _EXT_MAP.get(fmt, fmt)
    out = SS / f"process_data.{ext}"
    ok = canvas.export_file(str(out), fmt)
    size = out.stat().st_size if ok and out.exists() else 0
    check(f"export_{fmt}", ok and size > 0, f"ok={ok}, size={size}")
    if ok:
        print(f"    {name}: {size:,} bytes")

# Verify DOT content is flowchart (not calltree)
print("\n=== Content Verification ===")
dot_content = (SS / "process_data.dot").read_text(encoding="utf-8")
check("dot_is_flowchart", "Flowchart" in dot_content, "DOT should say 'Flowchart'")
check("dot_has_diamond", "diamond" in dot_content, "Should have diamond shape")
check("dot_has_trapezium", "trapezium" in dot_content, "Should have trapezium shape")
check("dot_has_true_false", "True" in dot_content and "False" in dot_content, "Should have True/False labels")

# Verify SVG content
svg_content = (SS / "process_data.svg").read_text(encoding="utf-8")
check("svg_has_nodes", "process_data" in svg_content, "SVG should contain function name")
check("svg_has_validate", "validate" in svg_content, "SVG should contain validate call")

# Verify HTML
html_content = (SS / "process_data.html").read_text(encoding="utf-8")
check("html_has_structure", "node" in html_content, "HTML should have node elements")

# Verify Excel
from openpyxl import load_workbook
wb = load_workbook(str(SS / "process_data.xlsx"))
ws = wb.active
check("excel_has_data", ws.max_row > 2, f"Excel has {ws.max_row} rows")

# Verify PNG is valid image
png_data = (SS / "process_data.png").read_bytes()
check("png_valid", png_data[:4] == b"\x89PNG", "PNG header invalid")

# Screenshot
win.grab().save(str(SS / "gui_flowchart_view.png"))

print("\n" + "=" * 50)
if issues:
    print(f"ISSUES: {len(issues)}")
    for i, issue in enumerate(issues, 1):
        print(f"  {i}. {issue}")
    sys.exit(1)
else:
    print("ALL FLOWCHART EXPORT CHECKS PASSED")
    sys.exit(0)
